"""
Bulk product import from .xlsx / .csv files.

  build_preview(file_bytes, filename, create_missing_cats, import_as_hidden)
    -> ImportPreview   (read-only DB lookups – no writes)

  execute_import(file_bytes, filename, options)
    -> result dict     (all writes inside one transaction)

  match_photos_zip(zip_bytes, image_file_map, replace_existing)
    -> photo result dict
"""
import csv
import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from typing import List, Optional

MAX_ROWS = 500

# Lowercase spreadsheet header -> canonical field name  (None = ignore silently)
HEADER_MAP = {
    "image":                None,           # embedded thumbnail column – ignored
    "sku":                  "sku",
    "product name":         "name",
    "name":                 "name",
    "category":             "category_name",
    "short description":    "summary",
    "description":          "summary",
    "age range":            "age_range",
    "cost price (pkr)":     None,           # internal – never stored
    "cost price":           None,
    "selling price (pkr)":  "selling_price",
    "selling price":        "selling_price",
    "sale price (pkr)":     "sale_price",
    "sale price":           "sale_price",
    "profit (pkr)":         None,           # internal – never stored
    "profit":               None,
    "margin %":             None,           # internal – never stored
    "margin":               None,
    "stock qty":            "stock",
    "stock":                "stock",
    "qty":                  "stock",
    "barcode":              None,           # internal – never stored
    "status":               "status",
    "image file":           "image_file",
}

# Category name (lowercase) -> Need names to assign ONLY to newly created products
CATEGORY_NEEDS_MAP = {
    "fine motor":         ["Strong little hands"],
    "sensory play":       ["Sensory exploring", "Calm and focus"],
    "gross motor":        ["Movement and balance"],
    "speech & language":  ["First words"],
    "thinking & puzzles": ["Thinking skills"],
    "social & emotional": ["Calm and focus"],
}


class ImportRowError(Exception):
    pass


@dataclass
class PreviewRow:
    row_num: int
    name: str
    sku: str
    category_name: str
    price: int
    compare_at_price: Optional[int]
    stock: int
    is_active: bool
    image_file: str
    action: str   # "create" | "update" | "error"
    error: str = ""


@dataclass
class ImportPreview:
    rows: List[PreviewRow] = field(default_factory=list)
    will_create: int = 0
    will_update: int = 0
    error_count: int = 0
    truncated: bool = False
    total_in_file: int = 0

    @property
    def preview_rows(self):
        return self.rows[:20]

    @property
    def error_rows(self):
        return [r for r in self.rows if r.action == "error"]


# ── Parsing helpers ────────────────────────────────────────────────────────────

def _parse_age(age_val):
    """
    Return (age_from, age_to).
    None / empty -> default (1, 6).
    Unparseable -> raise ValueError.
    """
    if age_val is None:
        return 1, 6
    s = str(age_val).strip()
    if s.endswith(".0"):
        s = s[:-2]  # strip trailing .0 from xlsx numeric cells
    if not s:
        return 1, 6
    # "6+" -> age_from=6, age_to=12
    m = re.match(r"^(\d+)\s*\+$", s)
    if m:
        return int(m.group(1)), 12
    # "X-Y" or "X–Y" or "X—Y"
    m = re.match(r"^(\d+)\s*[–\-—]\s*(\d+)$", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    # "X to Y"
    m = re.match(r"^(\d+)\s+to\s+(\d+)$", s, re.IGNORECASE)
    if m:
        return int(m.group(1)), int(m.group(2))
    raise ValueError(f"Cannot parse age range: {s!r}")


def _safe_int(val, field_name, required=False):
    """Parse val to int. Empty -> 0 (or raise if required=True). Invalid -> raise ImportRowError."""
    if val is None or str(val).strip().lower() in ("", "n/a", "-", "none", "null"):
        if required:
            raise ImportRowError(f"{field_name} is required")
        return 0
    try:
        return int(float(str(val).strip().replace(",", "").replace(" ", "")))
    except (ValueError, TypeError):
        raise ImportRowError(f"Invalid {field_name}: {val!r}")


def _normalize_row(raw_row):
    """Map a raw (lowercase-keyed) header dict to canonical field names."""
    canonical = {}
    for raw_key, raw_val in raw_row.items():
        mapped = HEADER_MAP.get(raw_key)
        if mapped is not None and mapped not in canonical:
            canonical[mapped] = raw_val
    return canonical


def _compute_pricing(canonical):
    """Return (price, compare_at_price) based on Sale Price / Selling Price logic."""
    selling = _safe_int(canonical.get("selling_price"), "Selling Price", required=True)
    sale_raw = canonical.get("sale_price")
    sale = _safe_int(sale_raw, "Sale Price") if sale_raw not in (None, "") else 0
    if sale and 0 < sale < selling:
        return sale, selling
    return selling, None


def _parse_xlsx(file_bytes):
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl is required to import .xlsx files. Run: pip install openpyxl"
        )
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = None
    for sheet_name in wb.sheetnames:
        if sheet_name.strip().lower() == "product catalog":
            ws = wb[sheet_name]
            break
    if ws is None:
        ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    raw_headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        result.append(dict(zip(raw_headers, row)))
    return result


def _parse_csv(file_bytes):
    text = file_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{k.strip().lower(): v for k, v in row.items()} for row in reader if any(row.values())]


def _parse_file(file_bytes, filename):
    if filename.lower().endswith(".csv"):
        return _parse_csv(file_bytes)
    return _parse_xlsx(file_bytes)


# ── Public API ─────────────────────────────────────────────────────────────────

def build_preview(file_bytes, filename, create_missing_categories=True, import_as_hidden=True):
    """Parse file, run read-only DB lookups, return ImportPreview. No writes."""
    from store.models import Category, Product

    raw_rows = _parse_file(file_bytes, filename)
    total_in_file = len(raw_rows)
    preview = ImportPreview(total_in_file=total_in_file)

    if len(raw_rows) > MAX_ROWS:
        raw_rows = raw_rows[:MAX_ROWS]
        preview.truncated = True

    cat_map = {c.name.lower(): c for c in Category.objects.all()}
    sku_map = {(p.sku or "").lower(): p for p in Product.objects.filter(sku__gt="")}
    name_map = {p.name.lower(): p for p in Product.objects.all()}

    for i, raw in enumerate(raw_rows):
        row_num = i + 2  # 1-indexed + header row
        canonical = _normalize_row(raw)
        name = str(canonical.get("name") or "").strip()
        sku = str(canonical.get("sku") or "").strip()

        if not name:
            preview.rows.append(PreviewRow(
                row_num=row_num, name="", sku=sku, category_name="",
                price=0, compare_at_price=None, stock=0, is_active=False,
                image_file="", action="error", error="Missing Product Name",
            ))
            preview.error_count += 1
            continue

        try:
            cat_name = str(canonical.get("category_name") or "").strip()
            cat_exists = bool(cat_map.get(cat_name.lower()))
            if cat_name and not cat_exists and not create_missing_categories:
                raise ImportRowError(
                    f"Category ‘{cat_name}’ not found. "
                    "Tick ‘Create missing categories’ to add it automatically."
                )

            age_from, age_to = _parse_age(canonical.get("age_range"))
            price, compare_at_price = _compute_pricing(canonical)

            stock = _safe_int(canonical.get("stock"), "Stock Qty")
            status_raw = str(canonical.get("status") or "").strip().lower()
            image_file = str(canonical.get("image_file") or "").strip()

            existing = (sku_map.get(sku.lower()) if sku else None) or name_map.get(name.lower())
            action = "update" if existing else "create"
            # For new products, respect import_as_hidden; for existing, leave is_active unchanged
            is_active = (existing.is_active if existing else not import_as_hidden)

            preview.rows.append(PreviewRow(
                row_num=row_num, name=name[:120], sku=sku, category_name=cat_name,
                price=price, compare_at_price=compare_at_price,
                stock=stock, is_active=is_active,
                image_file=image_file, action=action,
            ))
            if action == "create":
                preview.will_create += 1
            else:
                preview.will_update += 1

        except (ImportRowError, ValueError) as exc:
            preview.rows.append(PreviewRow(
                row_num=row_num, name=name, sku=sku, category_name="",
                price=0, compare_at_price=None, stock=0, is_active=False,
                image_file="", action="error", error=str(exc),
            ))
            preview.error_count += 1

    return preview


def execute_import(file_bytes, filename, options):
    """
    Parse file and import products inside one DB transaction.

    options:
        create_missing_categories  bool  (default True)
        import_as_hidden           bool  (default True)

    Returns:
        created        int
        updated        int
        errors         list[str]
        image_file_map dict{ str -> Product }   (keyed by image_file lowercase)
    """
    from django.db import transaction
    from django.utils.text import slugify
    from store.models import Category, Need, Product

    create_missing_cats = options.get("create_missing_categories", True)
    import_as_hidden = options.get("import_as_hidden", True)

    raw_rows = _parse_file(file_bytes, filename)
    if len(raw_rows) > MAX_ROWS:
        raw_rows = raw_rows[:MAX_ROWS]

    result = {"created": 0, "updated": 0, "errors": [], "image_file_map": {}}

    with transaction.atomic():
        cat_map = {c.name.lower(): c for c in Category.objects.all()}
        sku_map = {(p.sku or "").lower(): p for p in Product.objects.filter(sku__gt="")}
        name_map = {p.name.lower(): p for p in Product.objects.all()}
        need_map = {n.name.lower(): n for n in Need.objects.all()}

        for i, raw in enumerate(raw_rows):
            row_num = i + 2
            canonical = _normalize_row(raw)
            name = str(canonical.get("name") or "").strip()[:120]
            if not name:
                result["errors"].append(f"Row {row_num}: Missing Product Name")
                continue

            sku = str(canonical.get("sku") or "").strip()[:40]

            try:
                # ── Category ──────────────────────────────────────────────────
                cat_name = str(canonical.get("category_name") or "").strip()[:60]
                cat = cat_map.get(cat_name.lower()) if cat_name else None
                if not cat and cat_name:
                    if create_missing_cats:
                        cat = Category.objects.create(name=cat_name)
                        cat_map[cat_name.lower()] = cat
                    else:
                        raise ImportRowError(f"Category ‘{cat_name}’ not found")
                if not cat:
                    raise ImportRowError("Category is required")

                # ── Age ───────────────────────────────────────────────────────
                age_from, age_to = _parse_age(canonical.get("age_range"))

                # ── Pricing ───────────────────────────────────────────────────
                price, compare_at_price = _compute_pricing(canonical)

                # ── Other fields ──────────────────────────────────────────────
                stock = _safe_int(canonical.get("stock"), "Stock Qty")
                summary = str(canonical.get("summary") or "").strip()[:200]
                image_file = str(canonical.get("image_file") or "").strip()

                # ── Match existing product ────────────────────────────────────
                existing = (sku_map.get(sku.lower()) if sku else None) or name_map.get(name.lower())

                if existing:
                    existing.category = cat
                    existing.price = price
                    existing.compare_at_price = compare_at_price
                    existing.stock = stock
                    existing.age_from = age_from
                    existing.age_to = age_to
                    if summary:
                        existing.summary = summary
                    if sku:
                        existing.sku = sku
                    # Never touch: description, helps_with, in_the_box, needs,
                    #              is_active, is_featured, therapist_pick, image, slug
                    existing.save(update_fields=[
                        "category", "price", "compare_at_price",
                        "stock", "age_from", "age_to", "summary", "sku",
                    ])
                    if sku and sku.lower() not in sku_map:
                        sku_map[sku.lower()] = existing
                    result["updated"] += 1
                    product = existing

                else:
                    # Generate a unique slug
                    base_slug = slugify(name)[:120] or "product"
                    slug_val = base_slug
                    n_try = 2
                    while Product.objects.filter(slug=slug_val).exists():
                        slug_val = f"{base_slug}-{n_try}"
                        n_try += 1

                    product = Product.objects.create(
                        name=name,
                        slug=slug_val,
                        sku=sku,
                        category=cat,
                        price=price,
                        compare_at_price=compare_at_price,
                        stock=stock,
                        age_from=age_from,
                        age_to=age_to,
                        summary=summary or name[:200],
                        description="",
                        is_active=not import_as_hidden,
                        is_featured=False,
                        therapist_pick=False,
                    )

                    # Assign needs based on category (new products only)
                    for need_name in CATEGORY_NEEDS_MAP.get(cat_name.lower(), []):
                        need_obj = need_map.get(need_name.lower())
                        if need_obj:
                            product.needs.add(need_obj)

                    if sku:
                        sku_map[sku.lower()] = product
                    name_map[name.lower()] = product
                    result["created"] += 1

                if image_file:
                    result["image_file_map"][image_file.lower()] = product

            except (ImportRowError, ValueError) as exc:
                result["errors"].append(f"Row {row_num} ({name}): {exc}")

    return result


def match_photos_zip(zip_bytes, image_file_map, replace_existing=False):
    """
    Match ZIP entries to products by the 'Image File' column basename.

    image_file_map: { image_filename_lower: Product }

    Returns:
        attached  list[str]  product names that received a photo
        no_match  list[str]  ZIP filenames with no matching product
        no_photo  list[str]  products still without a main photo
        errors    list[str]
    """
    from store.bulk_upload import ALLOWED_EXTS
    from PIL import Image as PilImage
    from django.core.files.base import ContentFile

    results: dict = {"attached": [], "no_match": [], "no_photo": [], "errors": []}

    if not zip_bytes or not image_file_map:
        results["no_photo"] = [p.name for p in image_file_map.values() if not p.image]
        return results

    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            for entry in zf.namelist():
                parts = entry.replace("\\", "/").split("/")
                if any(p.startswith("__MACOSX") or (p.startswith(".") and p) for p in parts):
                    continue
                basename = parts[-1]
                if not basename or basename.startswith("."):
                    continue
                ext = os.path.splitext(basename)[1].lower().lstrip(".")
                if ext not in ALLOWED_EXTS:
                    continue

                product = image_file_map.get(basename.lower())
                if product is None:
                    if basename not in results["no_match"]:
                        results["no_match"].append(basename)
                    continue

                if product.image and not replace_existing:
                    continue

                img_bytes = zf.read(entry)
                try:
                    PilImage.open(io.BytesIO(img_bytes)).verify()
                except Exception as exc:
                    results["errors"].append(f"{basename}: invalid image ({exc})")
                    continue

                product.image.save(basename, ContentFile(img_bytes), save=True)
                if product.name not in results["attached"]:
                    results["attached"].append(product.name)

    except zipfile.BadZipFile as exc:
        results["errors"].append(f"Invalid ZIP file: {exc}")
    except Exception as exc:
        results["errors"].append(f"Photo import error: {exc}")

    for product in image_file_map.values():
        product.refresh_from_db(fields=["image"])
        if not product.image and product.name not in results["no_photo"]:
            results["no_photo"].append(product.name)

    return results
